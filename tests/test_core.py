"""Automated checks for Kannada Digitizer.  Run:  pip install pytest  then  pytest

No internet or API key is needed: AI and translation services are replaced with stand-ins.
"""
import io
import json
import os
import shutil
import sys
import types
from pathlib import Path

import pytest
from PIL import Image

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

import ai_tools      # noqa: E402
import exports       # noqa: E402
import llm           # noqa: E402
import ocr_pipeline as ocr   # noqa: E402
import storage       # noqa: E402

KANNADA = "ಕರ್ನಾಟಕ ಸರ್ಕಾರದ ಸೂಚನೆ\nಎಲ್ಲಾ ವಿದ್ಯಾರ್ಥಿಗಳಿಗೆ ಪರೀಕ್ಷೆಯು ನವೆಂಬರ್ ತಿಂಗಳಲ್ಲಿ ನಡೆಯಲಿದೆ."


@pytest.fixture(autouse=True)
def clean_env(monkeypatch):
    for k in ("GEMINI_API_KEY", "GOOGLE_API_KEY", "ANTHROPIC_API_KEY", "GEMINI_MODEL"):
        monkeypatch.delenv(k, raising=False)
    ai_tools._CACHE.clear()
    llm._state.update(models=None, good=None)


# ------------------------------------------------------------------ storage
@pytest.fixture()
def db(tmp_path, monkeypatch):
    monkeypatch.setattr(storage, "DB_PATH", tmp_path / "t.db")
    storage.init_db()
    return storage


def test_search_finds_kannada_and_english(db):
    i = db.add_document("notice", KANNADA, "tesseract", 95, 1, "The exam is in November")
    assert [h["id"] for h in db.search("ಪರೀಕ್ಷೆ")] == [i]
    hit = db.search("EXAM")[0]
    assert hit["id"] == i and hit["matched"] == "English translation"
    assert db.search("zzzzz") == []


def test_short_queries_fall_back_to_like(db):
    db.add_document("a", "ಅಮ್ಮ ಬಂದಳು", "t", 90, 1)
    assert len(db.search("ಅಮ್ಮ")) == 1


def test_update_and_delete_keep_index_in_sync(db):
    i = db.add_document("a", "ಹಳೆಯ ಪಠ್ಯ ಇಲ್ಲಿದೆ", "t", 90, 1, "old text here")
    db.update_document(i, text="ಹೊಸ ಪಠ್ಯ ಇಲ್ಲಿದೆ", translation="new text here")
    assert db.search("ಹಳೆಯ") == [] and db.search("old text") == []
    assert len(db.search("ಹೊಸ ಪಠ್ಯ")) == 1 and len(db.search("new text")) == 1
    db.delete_document(i)
    assert db.search("new text") == []


def test_old_database_is_upgraded(tmp_path, monkeypatch):
    import sqlite3
    path = tmp_path / "old.db"
    c = sqlite3.connect(path)
    c.execute("CREATE TABLE documents(id INTEGER PRIMARY KEY AUTOINCREMENT,title TEXT,text TEXT,"
              "engine TEXT,avg_conf REAL,pages INTEGER,created TEXT DEFAULT CURRENT_TIMESTAMP)")
    c.execute("CREATE VIRTUAL TABLE docs_fts USING fts5(title,text,content='documents',"
              "content_rowid='id',tokenize='trigram')")
    c.execute("INSERT INTO documents(title,text,engine,avg_conf,pages) VALUES('n',?,'t',90,1)", (KANNADA,))
    c.execute("INSERT INTO docs_fts(rowid,title,text) VALUES(1,'n',?)", (KANNADA,))
    c.commit(); c.close()
    monkeypatch.setattr(storage, "DB_PATH", path)
    storage.init_db()
    assert len(storage.search("ಪರೀಕ್ಷೆ")) == 1


# ---------------------------------------------------------------- AI layer
GOOD = json.dumps({"corrected": ["a", "b"], "pairs": ["one", "two"], "english": "One two."})


def test_ai_reply_parsing():
    assert ai_tools._parse_understanding(GOOD, 2)["english"] == "One two."
    assert ai_tools._parse_understanding("```json\n" + GOOD + "\n```", 2)
    assert ai_tools._parse_understanding(GOOD, 3) is None          # wrong line count
    assert ai_tools._parse_understanding("sorry, no", 2) is None


def test_understand_uses_ai_and_caches(monkeypatch):
    monkeypatch.setenv("GEMINI_API_KEY", "x")
    calls = []
    monkeypatch.setattr(llm, "complete", lambda *a, **k: calls.append(1) or GOOD)
    r = ai_tools.understand("line one\nline two")
    assert r["source"] == "ai" and r["lines"] == [("a", "one"), ("b", "two")]
    ai_tools.understand("line one\nline two")
    assert len(calls) == 1


def test_understand_falls_back_when_ai_fails(monkeypatch):
    monkeypatch.setenv("GEMINI_API_KEY", "x")
    def boom(*a, **k): raise llm.AIError("quota")
    monkeypatch.setattr(llm, "complete", boom)
    monkeypatch.setattr(ai_tools, "translate_free", lambda t, target="en": "free english")
    monkeypatch.setattr(ai_tools, "_free_per_line", lambda lines: ["x"] * len(lines))
    r = ai_tools.understand("line one\nline two")
    assert r["source"] == "free" and r["english"] == "free english"


def test_understand_never_raises_when_everything_fails(monkeypatch):
    monkeypatch.setattr(ai_tools, "translate_free", lambda t, target="en": "Translation failed. x")
    r = ai_tools.understand("line one")
    assert r["source"] == "none" and r["lines"] == []
    assert ai_tools.understand("   ")["source"] == "none"


def test_translation_joins_lines_into_sentences():
    seen = []
    out = ai_tools._translate_paragraphs("a\nb\n\nc", lambda t: seen.append(t) or t.upper(), 4000)
    assert seen == ["a b", "c"] and out == "A B\n\nC"


# --------------------------------------------------------- model fallback
def _fake_genai(monkeypatch, behaviour, names=("gemini-3.5-flash", "gemini-2.5-flash")):
    from google import genai
    calls = []

    class Models:
        def list(self):
            return [types.SimpleNamespace(name="models/" + n, supported_actions=["generateContent"])
                    for n in names]

        def generate_content(self, model, contents, config=None):
            calls.append(model)
            r = behaviour(model)
            if isinstance(r, Exception):
                raise r
            return types.SimpleNamespace(text=r)

    class Client:
        def __init__(self, api_key=None): self.models = Models()
    monkeypatch.setattr(genai, "Client", Client)
    monkeypatch.setenv("GEMINI_API_KEY", "SECRET-KEY")
    return calls


def test_gemini_moves_to_next_model_when_one_is_gone(monkeypatch):
    from google.genai import errors
    gone = errors.ClientError(404, {"error": {"message": "gone", "status": "NOT_FOUND"}})
    calls = _fake_genai(monkeypatch, lambda m: gone if m == "gemini-flash-latest" else "ok " + m)
    assert llm.complete("hi").startswith("ok ")
    assert calls[0] == "gemini-flash-latest" and len(calls) == 2
    assert llm._state["good"] == calls[-1]


def test_gemini_bad_key_stops_at_once_and_hides_the_key(monkeypatch):
    from google.genai import errors
    bad = errors.ClientError(400, {"error": {"message": "API key not valid SECRET-KEY",
                                              "status": "INVALID_ARGUMENT", "details": [{"reason": "API_KEY_INVALID"}]}})
    calls = _fake_genai(monkeypatch, lambda m: bad)
    with pytest.raises(llm.AIError) as e:
        llm.complete("hi")
    assert len(calls) == 1 and "SECRET-KEY" not in str(e.value)


def test_gemini_quota_everywhere_gives_a_friendly_error(monkeypatch):
    from google.genai import errors
    quota = errors.ClientError(429, {"error": {"message": "quota", "status": "RESOURCE_EXHAUSTED"}})
    _fake_genai(monkeypatch, lambda m: quota)
    with pytest.raises(llm.AIError) as e:
        llm.complete("hi")
    assert "quota" in str(e.value).lower()


def test_no_key_means_ai_is_off():
    assert not llm.available() and llm.provider() is None


# --------------------------------------------------------------- reading
def test_loader_rejects_bad_and_huge_files():
    with pytest.raises(ocr.UnreadableFile):
        ocr.load_pages(b"not an image", "x.png")
    with pytest.raises(ocr.UnreadableFile):
        ocr.load_pages(b"x" * (ocr.MAX_FILE_MB * 1024 * 1024 + 1), "x.png")


def test_loader_fixes_phone_rotation():
    im = Image.new("RGB", (400, 200), "white")
    ex = im.getexif(); ex[274] = 6
    buf = io.BytesIO(); im.save(buf, "JPEG", exif=ex.tobytes())
    assert ocr.load_pages(buf.getvalue(), "p.jpg")[0].size == (200, 400)


needs_tesseract = pytest.mark.skipif(shutil.which("tesseract") is None, reason="Tesseract not installed")


@needs_tesseract
def test_clean_notice_is_read_correctly():
    img = Image.open(ROOT / "samples" / "notice_tilted.png").convert("RGB")
    r = ocr.process_page(img, True, False)
    assert r.avg_conf > 85 and "ಪರೀಕ್ಷೆಯು" in r.text


@needs_tesseract
def test_white_on_dark_photo_is_read():
    img = Image.open(ROOT / "samples" / "hard_letter_a.png").convert("RGB")
    r = ocr.process_page(img, True, False)
    assert sum("\u0c80" <= ch <= "\u0cff" for ch in r.text) > 20


@needs_tesseract
def test_progress_messages_are_reported():
    msgs = []
    ocr.process_page(Image.open(ROOT / "samples" / "notice_tilted.png").convert("RGB"), True, False,
                     on_step=msgs.append)
    assert msgs and all(isinstance(m, str) for m in msgs)


# ----------------------------------------------------------------- exports
def test_word_export_has_english_and_table():
    from docx import Document
    d = Document(io.BytesIO(exports.to_docx("T", KANNADA, "Meaning", [("ಕ", "k")])))
    heads = [p.text for p in d.paragraphs if p.style.name.startswith("Heading")]
    assert heads == ["T", "English meaning", "Kannada text", "Line by line"] and len(d.tables) == 1


# ------------------------------------------------------ backend additions
def test_backup_restore_round_trip(tmp_path, monkeypatch):
    monkeypatch.setattr(storage, "DB_PATH", tmp_path / "a.db"); storage.init_db()
    storage.add_document("notice", KANNADA, "tesseract", 95, 1, "Exam in November")
    blob = storage.export_json()
    monkeypatch.setattr(storage, "DB_PATH", tmp_path / "b.db"); storage.init_db()
    assert storage.import_json(blob) == (1, 0)
    assert len(storage.search("exam")) == 1 and len(storage.search("ಪರೀಕ್ಷೆ")) == 1
    assert storage.import_json(blob) == (0, 1)                      # nothing is duplicated


@pytest.mark.parametrize("bad", [b"not json", b'{"app": "other", "documents": []}',
                                 b'{"app": "kannada-digitizer", "documents": "x"}'])
def test_restore_rejects_files_that_are_not_backups(tmp_path, monkeypatch, bad):
    monkeypatch.setattr(storage, "DB_PATH", tmp_path / "c.db"); storage.init_db()
    with pytest.raises(ValueError):
        storage.import_json(bad)


def test_database_uses_wal_for_concurrent_visitors(tmp_path, monkeypatch):
    monkeypatch.setattr(storage, "DB_PATH", tmp_path / "w.db"); storage.init_db()
    assert storage._conn().execute("PRAGMA journal_mode").fetchone()[0] == "wal"


def test_download_names_are_safe_but_keep_kannada():
    assert exports.safe_filename("../../etc/passwd") == "etc_passwd"
    assert exports.safe_filename("a/b\\c:d*?") == "a_b_c_d"
    assert exports.safe_filename("ಪರೀಕ್ಷೆ ಸೂಚನೆ") == "ಪರೀಕ್ಷೆ ಸೂಚನೆ"
    assert exports.safe_filename("") == "document" and len(exports.safe_filename("x" * 300)) <= 60


@needs_tesseract
def test_repeat_reads_are_served_from_memory():
    ocr._READ_CACHE.clear()
    data = (ROOT / "samples" / "notice_tilted.png").read_bytes()
    _, r1, hit1 = ocr.read_document(data, "n.png", True, False)
    _, r2, hit2 = ocr.read_document(data, "n.png", True, False)
    assert (hit1, hit2) == (False, True) and r1[0].text == r2[0].text
    _, _, hit3 = ocr.read_document(data, "n.png", False, False)       # other options: read afresh
    assert hit3 is False


# ------------------------------------------------------- start-up guard report
def test_guard_report_spots_zip_files_and_extra_folders(tmp_path, monkeypatch):
    import diagnostics
    (tmp_path / "ocr_pipeline.py").write_text("x = 1\n")                    # old file: no version stamp
    (tmp_path / "update-v7.zip").write_bytes(b"PK")
    (tmp_path / "update-v7").mkdir()
    monkeypatch.setattr(diagnostics, "HERE", tmp_path)
    msg = diagnostics.explain([("ocr_pipeline.py", 0, 7)], {"ocr_pipeline": 7})
    assert "OLD one" in msg and "update-v7.zip" in msg and "extra folder" in msg
    assert "Upload the newest copy" in msg


def test_guard_report_says_reboot_when_only_memory_is_old(tmp_path, monkeypatch):
    import diagnostics
    (tmp_path / "llm.py").write_text("MODULE_VERSION = 7\n")
    monkeypatch.setattr(diagnostics, "HERE", tmp_path)
    msg = diagnostics.explain([("llm.py", 0, 7)], {"llm": 7})
    assert "Rebooting the app fixes this" in msg and "OLD one" not in msg


# ---------------------------------------------------------------- AI reader honesty
def test_ai_reader_saying_no_text_is_treated_as_empty(monkeypatch):
    monkeypatch.setenv("GEMINI_API_KEY", "x")
    monkeypatch.setattr(llm, "complete", lambda *a, **k: "NO_TEXT")
    r = ocr.run_vision(Image.new("RGB", (40, 40), "white"))
    assert r.text == "" and r.engine == "vision"


def test_ai_reader_prompt_asks_for_no_text_marker(monkeypatch):
    seen = []
    monkeypatch.setenv("GEMINI_API_KEY", "x")
    monkeypatch.setattr(llm, "complete", lambda p, *a, **k: seen.append(p) or "ಅ")
    ocr.run_vision(Image.new("RGB", (40, 40), "white"))
    assert "NO_TEXT" in seen[0]

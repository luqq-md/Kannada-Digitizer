"""Export helpers: TXT, DOCX and searchable PDF."""
import io
import re
import unicodedata

import pytesseract
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt
from PIL import Image

from ocr_pipeline import TESSDATA_DIR


def to_txt(text: str) -> bytes:
    return text.encode("utf-8")


def to_docx(title: str, text: str, english: str = "", pairs=None) -> bytes:
    """Word file: the English meaning first (if we have it), then the Kannada text, then an
    optional line-by-line table so a reader can follow along."""
    doc = Document()
    doc.add_heading(title, level=1)

    def para(txt, kannada=True, size=12):
        p = doc.add_paragraph()
        run = p.add_run(txt)
        run.font.size = Pt(size)
        if kannada:
            run.font.name = "Noto Sans Kannada"
            run._element.rPr.rFonts.set(qn("w:cs"), "Noto Sans Kannada")   # complex-script font
        return p

    if english:
        doc.add_heading("English meaning", level=2)
        for block in english.split("\n"):
            if block.strip():
                para(block, kannada=False)
    doc.add_heading("Kannada text", level=2)
    for line in text.split("\n"):
        para(line)
    if pairs:
        doc.add_heading("Line by line", level=2)
        table = doc.add_table(rows=0, cols=2)
        table.style = "Light Grid Accent 1"
        for kn, en in pairs:
            cells = table.add_row().cells
            cells[0].text, cells[1].text = kn, en
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def to_searchable_pdf(images: list) -> bytes:
    """PDF that looks like the original scan but has an invisible text layer."""
    from pypdf import PdfWriter, PdfReader
    writer = PdfWriter()
    for img in images:
        pdf_bytes = pytesseract.image_to_pdf_or_hocr(
            img.convert("RGB"), lang="kan+eng", extension="pdf",
            config="--psm 6")
        for page in PdfReader(io.BytesIO(pdf_bytes)).pages:
            writer.add_page(page)
    buf = io.BytesIO()
    writer.write(buf)
    return buf.getvalue()


def safe_filename(title: str, fallback: str = "document") -> str:
    """A name that is safe to offer as a download: no slashes or odd symbols, but every
    letter, vowel sign and digit in any script (Kannada included) is kept."""
    keep = "".join(ch if (unicodedata.category(ch)[0] in "LMN" or ch in "-_. ") else "_"
                   for ch in (title or "").strip())
    name = re.sub(r"_+", "_", re.sub(r"\s+", " ", keep)).strip(" ._")[:60]
    return name or fallback
